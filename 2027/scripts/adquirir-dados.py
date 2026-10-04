"""Aquisição docente. Roteiros e avaliações usam os CSV locais, sem rede.

Uso: python scripts/adquirir-dados.py [--somente fmi|sidra|bcb|rtn|wdi]
Requer curl no PATH e openpyxl apenas para RTN. Mantém resposta bruta, URL e hash.
"""
from pathlib import Path
import argparse, csv, hashlib, json, shutil, subprocess, datetime, sys, time

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"
NOW = datetime.datetime.now(datetime.timezone.utc).isoformat()
COUNTRIES = ["BRA", "USA", "COL", "ARG", "CHL", "MEX", "PER", "URY", "CAN", "GBR", "DEU", "FRA", "ZAF", "IND"]
NAMES = dict(zip(COUNTRIES, ["Brasil", "Estados Unidos", "Colômbia", "Argentina", "Chile", "México", "Peru", "Uruguai", "Canadá", "Reino Unido", "Alemanha", "França", "África do Sul", "Índia"]))
FMI = ["NGDP_RPCH", "PCPIPCH", "GGXWDG_NGDP", "GGXCNL_NGDP", "BCA_NGDPD", "LUR", "NGDPDPC", "PPPGDP"]

def write_csv(name, rows):
    rows = list(rows)
    if not rows:
        raise ValueError(f"Nenhuma linha válida para {name}")
    target = OUT / f"{name}.csv"
    stage = target.with_suffix(".csv.tmp")
    with stage.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    stage.replace(target)
    return target

def fetch(url, name, binary=False):
    path = RAW / name
    stage = path.with_suffix(path.suffix + ".tmp")
    command = [shutil.which("curl.exe") or shutil.which("curl") or "curl", "--fail", "--location", "--retry", "2", "--max-time", "60", "--silent", "--show-error", url, "--output", str(stage)]
    subprocess.run(command, check=True)
    if not stage.exists() or stage.stat().st_size == 0:
        raise ValueError(f"Resposta vazia: {url}")
    content = stage.read_bytes()
    if not binary:
        json.loads(content)
    elif content[:2] != b"PK":
        raise ValueError("Planilha RTN não é XLSX (possível redirecionamento/HTML).")
    stage.replace(path)
    return path, content

def add_manifest(path, url, description, extra=""):
    manifests.append(dict(arquivo=str(path.relative_to(ROOT)).replace("\\", "/"), fonte=url,
      extraido_utc=NOW, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
      descricao=description, observacoes=extra))

def acquire_fmi():
    rows = []; metadata = []
    for code in FMI:
        url = f"https://www.imf.org/external/datamapper/api/v2/{code}"
        path, content = fetch(url, f"fmi-{code}.json")
        payload = json.loads(content)
        meta = payload.get("indicators", {}).get(code)
        if not meta or code not in payload.get("values", {}):
            raise ValueError(f"Indicador FMI não encontrado: {code}")
        projection = meta.get("projection-year")
        metadata.append(dict(indicador=code, nome=meta.get("label"), unidade=meta.get("unit"),
          fonte=meta.get("source"), inicio_projecao=projection, descricao=meta.get("description"),
          atualizado_fonte=meta.get("last-modified"), extraido_utc=NOW))
        values = payload["values"][code]
        for country in COUNTRIES:
            for year in range(2000, 2027):
                value = values.get(country, {}).get(str(year))
                status = "projeção" if projection and year >= int(projection) else "histórico ou estimativa; consultar WEO"
                rows.append(dict(pais=NAMES[country], iso3=country, ano=year, indicador=code,
                  valor=value, unidade=meta.get("unit"), status=status, edicao=meta.get("source")))
        add_manifest(path, url, f"FMI DataMapper {code}", "API pode retornar países adicionais; filtro local explícito. Histórico não implica observado definitivo.")
    p = write_csv("fmi_paises", rows); add_manifest(p, "https://www.imf.org/external/datamapper/api/", "Recorte de 14 países, 2000–2026, formato longo")
    p = write_csv("fmi_metadados", metadata); add_manifest(p, "https://www.imf.org/external/datamapper/api/v2/indicators", "Definições, unidade e início das projeções por indicador")

def sidra(name, table, variable, periods, classification=""):
    url = f"https://apisidra.ibge.gov.br/values/t/{table}/n1/all/v/{variable}/p/{periods}{classification}"
    path, content = fetch(url, f"sidra-{name}.json")
    items = json.loads(content)
    if not isinstance(items, list) or len(items) < 2:
        raise ValueError(f"Resposta SIDRA sem observações: {name}")
    labels = items[0]
    time_key = next(k[:-1] + "C" for k, v in labels.items() if k.endswith("N") and ("Mês" in v or "Trimestre" in v or v == "Ano"))
    rows = []
    for item in items[1:]:
        value = item.get("V")
        number = None if value in ("..", "...", "-", "X", "") else float(value)
        cat_key = "D4N" if "D4N" in item else "D3N"
        rows.append(dict(periodo=item[time_key], valor=number, unidade=item.get("MN"),
          variavel=item.get("D2N"), categoria=item.get(cat_key), tabela=table, codigo_variavel=variable))
    add_manifest(path, url, f"IBGE SIDRA tabela {table}, variável {variable}")
    p = write_csv(name, rows); add_manifest(p, url, f"Recorte SIDRA {name}")

def acquire_sidra():
    months = ",".join(f"{y}{m:02}" for y in range(2019, 2026) for m in range(1, 13))
    quarters = ",".join(f"{y}{q:02}" for y in range(2019, 2026) for q in range(1, 5))
    years = ",".join(str(y) for y in range(2012, 2025))
    sidra("ipca", 1737, 63, months)
    sidra("pib_volume", 1620, 583, quarters, "/c11255/90707,90687,90691,90696")
    sidra("desocupacao", 6381, 4099, months)
    sidra("rendimento_real", 6389, 5932, months)
    sidra("gini_oficial", 7435, 10681, years)

def acquire_bcb():
    rows = []
    definitions = {432: ("selic_meta", "% a.a.; meta diária"), 13762: ("dbgg", "% PIB; estoque mensal"),
      4513: ("dlsp", "% PIB; estoque mensal"), 5793: ("nfsp_primaria", "% PIB; fluxo acumulado em 12 meses; déficit positivo"),
      5760: ("juros_nominais", "% PIB; fluxo acumulado em 12 meses"), 5727: ("nfsp_nominal", "% PIB; fluxo acumulado em 12 meses; déficit positivo")}
    for code, (label, unit) in definitions.items():
        url=f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados?formato=json&dataInicial=01/01/2019&dataFinal=31/12/2025"
        path, raw=fetch(url, f"bcb-{code}.json")
        for item in json.loads(raw):
            date=datetime.datetime.strptime(item["data"], "%d/%m/%Y").date().isoformat()
            rows.append(dict(data=date, indicador=label, codigo=code, valor=float(item["valor"]), unidade=unit))
        add_manifest(path,url,f"BCB SGS {code}")
    p=write_csv("bcb", rows); add_manifest(p,"https://www3.bcb.gov.br/sgspub/", "BCB, códigos documentados")

def acquire_wdi():
    rows=[]
    codes={"SI.POV.GINI":"Índice de Gini, escala 0–100", "SL.UEM.TOTL.ZS":"Desemprego, estimativa modelada OIT, % força de trabalho"}
    countries=";".join(COUNTRIES)
    for code,unit in codes.items():
        url=f"https://api.worldbank.org/v2/country/{countries}/indicator/{code}?format=json&date=2000:2025&per_page=20000"
        path,raw=fetch(url,f"wdi-{code}.json")
        payload=json.loads(raw)
        if not isinstance(payload,list) or len(payload)!=2 or not payload[1]: raise ValueError(f"WDI inválido: {code}")
        for item in payload[1]:
            iso=item["countryiso3code"]
            rows.append(dict(pais=NAMES.get(iso,iso),iso3=iso,ano=int(item["date"]),indicador=code,valor=item["value"],unidade=unit,status="estimativa; conferir metadados WDI",edicao=f"WDI extraído {NOW[:10]}"))
        add_manifest(path,url,f"WDI {code}")
    p=write_csv("wdi_complemento",rows); add_manifest(p,"https://data.worldbank.org/indicator", "Complemento desigualdade e emprego")

def acquire_rtn():
    import openpyxl
    url="https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show?id=resultado-do-tesouro-nacional"
    path,raw=fetch(url,"rtn-catalogo.json")
    resources=json.loads(raw)["result"]["resources"]
    resource=next(r for r in resources if r["id"]=="527ccdb1-3059-42f3-bf23-b5e3ab4c6dc6")
    path,raw=fetch(resource["url"],"rtn-serie-historica.xlsx",binary=True)
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    # Selecionar aba pela indicação nominal/corrente, nunca pela posição.
    matches=[]
    for sheet_candidate in wb.worksheets:
        heading=" ".join(str(v) for row in sheet_candidate.iter_rows(max_row=4,values_only=True) for v in row if v is not None).lower()
        if "resumida" in heading and "valores correntes" in heading and "mensal" in heading:
            matches.append(sheet_candidate.title)
    if len(matches)!=1: raise ValueError(f"Identifique a aba nominal antes de extrair RTN: {wb.sheetnames}")
    sheet=wb[matches[0]]; cells=list(sheet.iter_rows(values_only=True))
    dates={}
    for row in cells[:15]:
        for i,value in enumerate(row):
            if isinstance(value,datetime.datetime) and 2019<=value.year<=2025: dates[i]=value.date().isoformat()
    if not dates: raise ValueError("Cabeçalho mensal RTN não identificado; inspecione a planilha.")
    rows=[]
    for row in cells:
        label=next((v for v in row[:3] if isinstance(v,str) and v.strip()), "")
        if any(s in label.upper() for s in ["RECEITA LÍQUIDA", "DESPESA TOTAL", "RESULTADO PRIMÁRIO GOVERNO CENTRAL", "BENEFÍCIOS PREVIDENCIÁRIOS"]):
            for i,date in dates.items():
                if i<len(row) and isinstance(row[i],(int,float)):
                    rows.append(dict(data=date,rubrica=label,valor=row[i],unidade="R$ milhões correntes",aba=matches[0]))
    if not rows: raise ValueError("Rubricas RTN não identificadas.")
    add_manifest(path,resource["url"],"RTN planilha fonte", f"Aba {matches[0]}; extração por nome da rubrica")
    p=write_csv("rtn",rows); add_manifest(p,resource["url"],"RTN, rubricas selecionadas e valores nominais")

if __name__=="__main__":
    parser=argparse.ArgumentParser(); parser.add_argument("--somente",choices=["fmi","sidra","bcb","rtn","wdi"]); args=parser.parse_args()
    RAW.mkdir(parents=True,exist_ok=True); OUT.mkdir(parents=True,exist_ok=True)
    manifest_file=ROOT/"data"/"inventario-bases.json"
    manifests=json.loads(manifest_file.read_text(encoding="utf-8")) if manifest_file.exists() else []
    errors=[]
    for name,fn in [("fmi",acquire_fmi),("sidra",acquire_sidra),("bcb",acquire_bcb),("rtn",acquire_rtn),("wdi",acquire_wdi)]:
        if args.somente and name!=args.somente: continue
        try: fn(); print(f"OK: {name}",flush=True)
        except Exception as e: errors.append(dict(fonte=name,erro=str(e))); print(f"FALHA: {name}: {e}",flush=True)
    # Uma entrada por arquivo: aquisições posteriores substituem seu registro.
    manifest_file.write_text(json.dumps(list({x['arquivo']:x for x in manifests}.values()),ensure_ascii=False,indent=2),encoding="utf-8")
    (ROOT/"data"/"status-aquisicao.json").write_text(json.dumps(dict(executado_utc=NOW,falhas=errors),ensure_ascii=False,indent=2),encoding="utf-8")
    sys.exit(1 if errors else 0)
